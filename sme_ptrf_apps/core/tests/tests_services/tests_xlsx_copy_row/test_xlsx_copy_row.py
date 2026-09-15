from openpyxl import Workbook
from openpyxl.styles import Font

from sme_ptrf_apps.core.services.xlsx_copy_row import copy_row, copy_row_direct, insert_row


def _sheet_with_values(rows):
    wb = Workbook()
    ws = wb.active
    for row_idx, value in enumerate(rows, start=1):
        ws.cell(row=row_idx, column=1, value=value)
    return wb, ws


def test_copy_row_copia_valor_e_estilo():
    wb, ws = _sheet_with_values(["R1", "R2"])
    ws['A1'].font = Font(bold=True)

    copy_row(ws, source_row=1, dest_row=1, copy_data=True, copy_style=True)

    assert ws['A2'].value == "R1"
    assert ws['A2'].font.bold is True


def test_copy_row_sem_copy_data_nao_copia_valor():
    wb, ws = _sheet_with_values(["R1", "R2"])

    copy_row(ws, source_row=1, dest_row=1, copy_data=False, copy_style=True)

    assert ws['A2'].value == "R2"


def test_copy_row_sem_copy_style_nao_copia_estilo():
    wb, ws = _sheet_with_values(["R1", "R2"])
    ws['A1'].font = Font(bold=True)

    copy_row(ws, source_row=1, dest_row=1, copy_data=True, copy_style=False)

    assert ws['A2'].font.bold is not True


def test_copy_row_cascata_para_varias_linhas():
    wb, ws = _sheet_with_values(["R1", "R2", "R3"])

    copy_row(ws, source_row=1, dest_row=2, copy_data=True)

    # Copia em cascata: linha 2 recebe o valor da linha 1, e linha 3 recebe
    # o valor já copiado para a linha 2 (ambas acabam iguais à linha 1).
    assert ws['A2'].value == "R1"
    assert ws['A3'].value == "R1"


def test_copy_row_copia_merge_de_coluna():
    wb, ws = _sheet_with_values(["R1", "R2"])
    ws.merge_cells('A1:B1')

    copy_row(ws, source_row=1, dest_row=1, copy_data=True, copy_merged_columns=True)

    assert 'A2:B2' in [str(r) for r in ws.merged_cells.ranges]


def test_copy_row_sem_copy_merged_columns_nao_copia_merge():
    wb, ws = _sheet_with_values(["R1", "R2"])
    ws.merge_cells('A1:B1')

    copy_row(ws, source_row=1, dest_row=1, copy_data=True, copy_merged_columns=False)

    assert 'A2:B2' not in [str(r) for r in ws.merged_cells.ranges]


def test_copy_row_desfaz_merge_existente_no_destino():
    wb, ws = _sheet_with_values(["R1", "R2", "R3"])
    ws.merge_cells('A2:B2')

    # Não deve levantar exceção ao desfazer o merge que já existia na linha de destino.
    copy_row(ws, source_row=1, dest_row=1, copy_data=True)

    assert 'A2:B2' not in [str(r) for r in ws.merged_cells.ranges]


def test_insert_row_desloca_linhas_para_baixo():
    wb, ws = _sheet_with_values(["R1", "R2", "R3"])

    insert_row(ws, last_row=3, source_row=1)

    assert ws['A1'].value == "R1"
    assert ws['A2'].value == "R2"
    assert ws['A3'].value == "R2"
    assert ws['A4'].value == "R3"


def test_copy_row_direct_copia_valor_e_estilo_para_linha_especifica():
    wb, ws = _sheet_with_values(["R1", "R2", "R3"])
    ws['A1'].font = Font(bold=True)

    copy_row_direct(ws, source_row=1, dest_row=3, copy_data=True, copy_style=True)

    assert ws['A3'].value == "R1"
    assert ws['A3'].font.bold is True


def test_copy_row_direct_sem_copy_data_nao_copia_valor():
    wb, ws = _sheet_with_values(["R1", "R2", "R3"])

    copy_row_direct(ws, source_row=1, dest_row=3, copy_data=False, copy_style=True)

    assert ws['A3'].value == "R3"


def test_copy_row_direct_copia_merge_de_coluna_para_linha_especifica():
    wb, ws = _sheet_with_values(["R1", "R2", "R3"])
    ws.merge_cells('A1:B1')

    copy_row_direct(ws, source_row=1, dest_row=3, copy_data=True, copy_merged_columns=True)

    assert 'A3:B3' in [str(r) for r in ws.merged_cells.ranges]


def test_copy_row_direct_sem_copy_merged_columns_nao_copia_merge():
    wb, ws = _sheet_with_values(["R1", "R2", "R3"])
    ws.merge_cells('A1:B1')

    copy_row_direct(ws, source_row=1, dest_row=3, copy_data=True, copy_merged_columns=False)

    assert 'A3:B3' not in [str(r) for r in ws.merged_cells.ranges]


def test_copy_row_direct_desfaz_merge_existente_logo_apos_a_origem():
    wb, ws = _sheet_with_values(["R1", "R2", "R3"])
    ws.merge_cells('A2:B2')

    # Não deve levantar exceção ao desfazer o merge que já existia logo após a linha de origem.
    copy_row_direct(ws, source_row=1, dest_row=3, copy_data=True)

    assert 'A2:B2' not in [str(r) for r in ws.merged_cells.ranges]
